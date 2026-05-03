"""
export_customer_report.py
=========================
Generate a personalised buying recommendation report for a specific customer.
Exports as a professional Excel workbook they can act on immediately.

Usage:
  python3 export_customer_report.py --customer "BELGAL TRADE"
  python3 export_customer_report.py --customer "Next Retail Ltd"
  python3 export_customer_report.py --list   # list all customers
"""

import argparse
import sys
import os
from datetime import datetime

os.system("pip install openpyxl google-cloud-bigquery db-dtypes pyarrow --break-system-packages -q")

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from google.cloud import bigquery

PROJECT = "inventory-planning-model"
DATASET = "Model"
T = lambda n: f"`{PROJECT}`.`{DATASET}`.`{n}`"

parser = argparse.ArgumentParser()
parser.add_argument("--customer", default="", help="Customer name (exact match)")
parser.add_argument("--list", action="store_true", help="List all customers")
args = parser.parse_args()

client = bigquery.Client(project=PROJECT)

# ── List customers ────────────────────────────────────────────────
if args.list:
    df = client.query(f"SELECT DISTINCT customer FROM {T('buy_plan')} ORDER BY customer").to_dataframe()
    print("\nAll customers in buy_plan:")
    for c in df["customer"].tolist(): print(" ", c)
    sys.exit(0)

if not args.customer:
    print("Please provide --customer 'Customer Name' or use --list to see all customers")
    sys.exit(1)

cust = args.customer.strip()
print(f"\nGenerating report for: {cust}")

def bq(sql):
    try:
        return client.query(sql).to_dataframe()
    except Exception as e:
        print(f"  Query failed: {e}")
        return pd.DataFrame()

# ── Pull all data for this customer ──────────────────────────────
print("  Loading buy plan…")
buy = bq(f"""
    SELECT brand, style, colour, size,
           SUBSTR(year_month,1,7) AS month,
           ROUND(forecast_p50,0)  AS p50,
           ROUND(forecast_p90,0)  AS p90,
           ROUND(buy_qty,0)       AS buy_qty,
           moq,
           CAST(new_style_flag AS INT64)     AS new_style,
           CAST(below_moq_flag AS INT64)     AS below_moq,
           CAST(zero_forecast_flag AS INT64) AS zero_fcst
    FROM {T('buy_plan')}
    WHERE UPPER(customer) = UPPER('{cust}')
      AND buy_qty > 0
    ORDER BY brand, style, colour, size, month
""")

print("  Loading reorder alerts…")
reorder = bq(f"""
    SELECT brand, last_order_month, months_overdue,
           avg_order_qty, priority
    FROM {T('recommendations_reorder')}
    WHERE UPPER(customer) = UPPER('{cust}')
    ORDER BY months_overdue DESC
""")

print("  Loading cross-sell recommendations…")
crosssell = bq(f"""
    SELECT recommended_brand AS brand,
           ROUND(confidence_score,1) AS confidence,
           similar_customers_buying,
           priority
    FROM {T('recommendations_crosssell')}
    WHERE UPPER(customer) = UPPER('{cust}')
    ORDER BY confidence_score DESC
""")

print("  Loading new style recommendations…")
newstyle = bq(f"""
    SELECT brand, recommended_style AS style,
           ROUND(recommendation_score,1) AS score,
           ROUND(colour_compatibility*100,0) AS colour_match_pct,
           priority
    FROM {T('recommendations_newstyle')}
    WHERE UPPER(customer) = UPPER('{cust}')
      AND priority = 'HIGH'
    ORDER BY recommendation_score DESC
    LIMIT 50
""")

if buy.empty and reorder.empty and crosssell.empty:
    print(f"\nNo data found for '{cust}'. Try --list to see exact customer names.")
    sys.exit(1)

# ── Setup workbook ────────────────────────────────────────────────
wb = Workbook()

NAVY  = "0A1628"
TEAL  = "0D7377"
GOLD  = "D4A017"
LIGHT = "EEF4FB"
WHITE = "FFFFFF"

def hfill(c): return PatternFill("solid", fgColor=c)
def hfont(sz=10, bold=False, color=WHITE):
    return Font(name="Arial", size=sz, bold=bold, color=color)
def bfont(sz=10, color="000000"):
    return Font(name="Arial", size=sz, color=color)
thin_border = Border(
    left=Side(style='thin',color="C5D5E8"),right=Side(style='thin',color="C5D5E8"),
    top=Side(style='thin',color="C5D5E8"),bottom=Side(style='thin',color="C5D5E8"))

def style_cell(cell, fill_color=None, font=None, align="left", bold=False, sz=10, fc=None):
    if fill_color: cell.fill = hfill(fill_color)
    if fc:
        cell.font = Font(name="Arial",size=sz,bold=bold,color=fc)
    else:
        cell.font = Font(name="Arial",size=sz,bold=bold)
    cell.alignment = Alignment(horizontal=align,vertical="center",wrap_text=True)
    cell.border = thin_border

def make_header(ws, title, subtitle=""):
    ws.merge_cells(f"A1:H1")
    ws["A1"] = "NORDIC KOLEKTIVE VENTURES · DEMAND INTELLIGENCE"
    ws["A1"].fill = hfill(NAVY)
    ws["A1"].font = hfont(13, True)
    ws["A1"].alignment = Alignment(horizontal="center",vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Customer: {cust}  |  Report: {title}  |  {datetime.now().strftime('%d %b %Y')}"
    ws["A2"].fill = hfill(TEAL)
    ws["A2"].font = hfont(10)
    ws["A2"].alignment = Alignment(horizontal="center",vertical="center")
    ws.row_dimensions[2].height = 18

    if subtitle:
        ws.merge_cells("A3:H3")
        ws["A3"] = subtitle
        ws["A3"].fill = hfill(LIGHT)
        ws["A3"].font = Font(name="Arial",size=10,color=NAVY)
        ws["A3"].alignment = Alignment(horizontal="center",vertical="center")
        ws.row_dimensions[3].height = 16
        ws.append([])
    else:
        ws.append([])
        ws.append([])

# ══════════════════════════════════════════════════════════════════
# SHEET 1: BUY PLAN
# ══════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "1. Buy Plan"
make_header(ws1, "Recommended Buy Plan",
            f"AI-generated purchase recommendations · {buy['month'].nunique() if not buy.empty else 0} months · {buy['brand'].nunique() if not buy.empty else 0} brands")

if not buy.empty:
    headers = ["Brand","Style","Colour","Size","Month","P50 Forecast","P90 Forecast","Recommended Buy Qty","MOQ","New Style","Below MOQ"]
    ws1.append(headers)
    hr = ws1.max_row
    for i, h in enumerate(headers, 1):
        c = ws1.cell(row=hr, column=i)
        c.value = h; c.fill = hfill(NAVY)
        c.font = hfont(10, True); c.border = thin_border
        c.alignment = Alignment(horizontal="center",vertical="center")
    ws1.row_dimensions[hr].height = 20

    for ri, row in buy.iterrows():
        vals = [row["brand"],row["style"],row["colour"],row["size"],row["month"],
                int(row["p50"] or 0), int(row["p90"] or 0), int(row["buy_qty"] or 0),
                int(row["moq"] or 0),
                "⚠ New" if row["new_style"] else "", "⚠ Review" if row["below_moq"] else ""]
        ws1.append(vals)
        r = ws1.max_row
        fill = LIGHT if r%2==0 else WHITE
        for ci, v in enumerate(vals, 1):
            cell = ws1.cell(row=r, column=ci)
            cell.fill = hfill(fill); cell.font = bfont(10)
            cell.alignment = Alignment(vertical="center"); cell.border = thin_border
        if row["new_style"]: ws1.cell(row=r, column=10).fill = hfill("E6F1FB")
        if row["below_moq"]: ws1.cell(row=r, column=11).fill = hfill("FAEEDA")

    for i, w in enumerate([18,14,16,6,10,14,14,16,8,10,11], 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = f"A{hr+1}"

    # Brand summary box
    ws1.append([]); ws1.append([])
    r = ws1.max_row
    ws1.cell(row=r, column=1).value = "BRAND SUMMARY"
    ws1.cell(row=r, column=1).fill = hfill(TEAL)
    ws1.cell(row=r, column=1).font = hfont(10,True)
    ws1.merge_cells(f"A{r}:E{r}")

    ws1.append(["Brand","Months","Total P50","Total Buy Qty","Buy/P50 Ratio"])
    hr2 = ws1.max_row
    for i in range(1,6):
        c = ws1.cell(row=hr2,column=i)
        c.fill = hfill(NAVY); c.font = hfont(10,True); c.border = thin_border

    bs = buy.groupby("brand").agg(
        months=("month","nunique"), p50=("p50","sum"), buy=("buy_qty","sum")
    ).reset_index().sort_values("buy", ascending=False)
    bs["ratio"] = (bs["buy"]/bs["p50"].replace(0,1)).round(2)

    for _, row in bs.iterrows():
        ws1.append([row["brand"], int(row["months"]), int(row["p50"]),
                    int(row["buy"]), float(row["ratio"])])
        r = ws1.max_row
        fill = LIGHT if r%2==0 else WHITE
        for ci in range(1,6):
            c = ws1.cell(row=r,column=ci)
            c.fill = hfill(fill); c.font = bfont(10); c.border = thin_border
        ratio = float(row["ratio"])
        rc = ws1.cell(row=r, column=5)
        if 1<=ratio<=2.5: rc.fill = hfill("EAF3DE")
        elif ratio>2.5: rc.fill = hfill("FAEEDA")
        else: rc.fill = hfill("FCEBEB")
else:
    ws1.append(["No buy plan data found for this customer."])

# ══════════════════════════════════════════════════════════════════
# SHEET 2: REORDER ALERTS
# ══════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("2. Reorder Alerts")
make_header(ws2, "Reorder Timing Alerts",
            "Brands where your ordering is overdue based on your historical purchase intervals")

if not reorder.empty:
    headers = ["Brand","Last Ordered","Months Overdue","Avg Order Qty","Priority","Action Required"]
    ws2.append(headers)
    hr = ws2.max_row
    for i, h in enumerate(headers, 1):
        c = ws2.cell(row=hr, column=i)
        c.value = h; c.fill = hfill(NAVY)
        c.font = hfont(10,True); c.border = thin_border
        c.alignment = Alignment(horizontal="center",vertical="center")

    for _, row in reorder.iterrows():
        action = ("Call your rep immediately" if row["priority"]=="URGENT"
                  else "Schedule order this month")
        vals = [row["brand"], row.get("last_order_month","—"),
                float(row.get("months_overdue",0)),
                float(row.get("avg_order_qty",0)),
                row["priority"], action]
        ws2.append(vals)
        r = ws2.max_row
        fill = LIGHT if r%2==0 else WHITE
        for ci, v in enumerate(vals, 1):
            c = ws2.cell(row=r,column=ci)
            c.fill = hfill(fill); c.font = bfont(10)
            c.alignment = Alignment(vertical="center"); c.border = thin_border
        # Priority colouring
        pc = ws2.cell(row=r, column=5)
        if row["priority"] == "URGENT": pc.fill = hfill("FCEBEB"); pc.font = bfont(10,"A32D2D")
        else: pc.fill = hfill("FAEEDA"); pc.font = bfont(10,"854F0B")

    for i, w in enumerate([22,16,15,14,12,30], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
else:
    ws2.append(["No overdue reorders for this customer — all up to date!"])
    ws2.cell(row=ws2.max_row,column=1).font = Font(name="Arial",size=11,color="3B6D11")

# ══════════════════════════════════════════════════════════════════
# SHEET 3: CROSS-SELL
# ══════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("3. Cross-sell Opps")
make_header(ws3, "Cross-Sell Opportunities",
            "Brands that similar customers are buying — high-confidence opportunities for this account")

if not crosssell.empty:
    headers = ["Recommended Brand","Confidence Score","Similar Customers Buying","Priority","Why This Brand"]
    ws3.append(headers)
    hr = ws3.max_row
    for i, h in enumerate(headers, 1):
        c = ws3.cell(row=hr, column=i)
        c.value = h; c.fill = hfill(NAVY)
        c.font = hfont(10,True); c.border = thin_border
        c.alignment = Alignment(horizontal="center",vertical="center")

    for _, row in crosssell.iterrows():
        why = (f"{row.get('similar_customers_buying',0)} customers with similar "
               f"buying patterns already stock this brand")
        vals = [row["brand"], float(row.get("confidence",0)),
                int(row.get("similar_customers_buying",0)),
                row["priority"], why]
        ws3.append(vals)
        r = ws3.max_row
        fill = LIGHT if r%2==0 else WHITE
        for ci, v in enumerate(vals, 1):
            c = ws3.cell(row=r,column=ci)
            c.fill = hfill(fill); c.font = bfont(10)
            c.alignment = Alignment(vertical="center",wrap_text=True); c.border = thin_border

    for i, w in enumerate([24,16,20,12,45], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
else:
    ws3.append(["No cross-sell opportunities identified for this customer."])

# ══════════════════════════════════════════════════════════════════
# SHEET 4: NEW STYLES
# ══════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("4. New Styles")
make_header(ws4, "New Style Recommendations",
            "Styles not yet in your range — matched on colour compatibility, price alignment and market demand")

if not newstyle.empty:
    headers = ["Brand","Recommended Style","Match Score","Colour Compatibility","Priority","Why This Style"]
    ws4.append(headers)
    hr = ws4.max_row
    for i, h in enumerate(headers, 1):
        c = ws4.cell(row=hr, column=i)
        c.value = h; c.fill = hfill(NAVY)
        c.font = hfont(10,True); c.border = thin_border
        c.alignment = Alignment(horizontal="center",vertical="center")

    for _, row in newstyle.iterrows():
        why = (f"Colour profile {row.get('colour_match_pct',0):.0f}% compatible with your "
               f"existing range. Score {row.get('score',0):.1f}/10 based on market volume and price fit.")
        vals = [row["brand"], row["style"],
                float(row.get("score",0)),
                f"{row.get('colour_match_pct',0):.0f}%",
                row["priority"], why]
        ws4.append(vals)
        r = ws4.max_row
        fill = LIGHT if r%2==0 else WHITE
        for ci, v in enumerate(vals, 1):
            c = ws4.cell(row=r,column=ci)
            c.fill = hfill(fill); c.font = bfont(10)
            c.alignment = Alignment(vertical="center",wrap_text=True); c.border = thin_border
        # Score colouring
        score = float(row.get("score",0))
        sc = ws4.cell(row=r, column=3)
        if score>=9: sc.fill = hfill("EAF3DE")
        elif score>=7: sc.fill = hfill("E6F1FB")

    for i, w in enumerate([18,16,13,18,12,50], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
else:
    ws4.append(["No new style recommendations for this customer."])

# ── Save ──────────────────────────────────────────────────────────
safe_name = cust.replace(" ","_").replace("/","_").replace("'","")[:40]
stamp = datetime.now().strftime("%Y%m%d")
fname = f"Report_{safe_name}_{stamp}.xlsx"

wb.save(fname)
print(f"\n✓ Report saved: {fname}")
print(f"  Buy plan lines: {len(buy):,}")
print(f"  Reorder alerts: {len(reorder)}")
print(f"  Cross-sell opps: {len(crosssell)}")
print(f"  New style recs: {len(newstyle)}")
print(f"\nDownload:")
print(f"  cloudshell download ~/{fname}")
print(f"\nOr email directly to customer as attachment.")
